"""Builds the submission documents from tools/thread_content.py.

Run: python3 tools/build_docs.py
Output:
  entry/THREAD.md                     plain-text thread, ready to copy into X
  entry/Redbelly_Thread.docx          the same thread as a document, images inline
  entry/Redbelly_Thread.pdf           PDF of the .docx (needs LibreOffice; skipped if absent)
"""

import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from check_entry import x_length  # noqa: E402
from thread_content import POSTS, TITLE  # noqa: E402

ENTRY = ROOT / "entry"
IMAGES = ENTRY / "images"
RED = RGBColor(0xFF, 0x50, 0x50)
SLATE = RGBColor(0x28, 0x3A, 0x46)
GREY = RGBColor(0x84, 0x8B, 0x91)

INTRO = (
    "Format: X thread, {n} posts, {k} original images. Each post is under X's "
    "280-character limit. Post them in order as one thread, attach each image to the post "
    "marked below, and paste the alt text into the image description."
)

SOURCES_SHORT = [
    ("Redbelly Network home page", "https://redbelly.network"),
    ("Redbelly About us and timeline", "https://redbelly.network/about-us"),
    ("Redbelly asset tokenisation page", "https://redbelly.network/asset-tokenisation"),
    ("Redbelly article: What sets Redbelly Network apart",
     "https://redbelly.network/blog/what-sets-redbelly-network-apart"),
    ("Redbelly Project Acacia page", "https://redbelly.network/projectacacia"),
    ("RBA and DFCRC, Project Acacia final report, May 2026",
     "https://www.rba.gov.au/payments-and-infrastructure/central-bank-digital-currency/pdf/project-acacia-final-report.pdf"),
    ("RBA media release, 18 May 2026", "https://www.rba.gov.au/media-releases/2026/mr-26-13.html"),
    ("Tokeniser home page", "https://tokeniser.com"),
]


def build_markdown():
    n_img = sum(1 for p in POSTS if p.get("image"))
    out = [f"# {TITLE}", "", INTRO.format(n=len(POSTS), k=n_img), ""]
    for i, p in enumerate(POSTS, 1):
        out += [f"## {i}/{len(POSTS)}", "", "```text", p["text"], "```", ""]
        meta = f"*{x_length(p['text'])} / 280 characters*"
        if p.get("image"):
            out += [f"![Image for post {i}](images/{p['image']})", "",
                    f"**Attach:** `images/{p['image']}`  ", f"**Alt text:** {p['alt']}  "]
        out += [meta, ""]
    out += ["---", "", "## Sources", ""]
    out += [f"- {name}: {url}" for name, url in SOURCES_SHORT]
    out += ["", "Sentence-by-sentence source mapping: `FACT_CHECK.md`. Full source notes: `SOURCES.md`.", ""]
    (ENTRY / "THREAD.md").write_text("\n".join(out), encoding="utf8")
    print("wrote entry/THREAD.md")


def build_docx():
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = TITLE
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)

    h = doc.add_paragraph()
    r = h.add_run(TITLE)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(22), SLATE
    n_img = sum(1 for p in POSTS if p.get("image"))
    intro = doc.add_paragraph(INTRO.format(n=len(POSTS), k=n_img))
    intro.runs[0].font.color.rgb = GREY

    for i, p in enumerate(POSTS, 1):
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(14)
        hr = hp.add_run(f"{i}/{len(POSTS)}")
        hr.bold, hr.font.size, hr.font.color.rgb = True, Pt(14), RED
        for para in p["text"].split("\n\n"):
            bp = doc.add_paragraph(para)
            bp.paragraph_format.space_after = Pt(4)
        if p.get("image"):
            doc.add_picture(str(IMAGES / p["image"]), width=Cm(15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            ap = doc.add_paragraph()
            a1 = ap.add_run("Attach: ")
            a1.bold = True
            ap.add_run(f"images/{p['image']}")
            ap2 = doc.add_paragraph()
            a2 = ap2.add_run("Alt text: ")
            a2.bold = True
            ap2.add_run(p["alt"])
        cp = doc.add_paragraph(f"{x_length(p['text'])} / 280 characters")
        cp.runs[0].font.size, cp.runs[0].font.color.rgb = Pt(9), GREY

    doc.add_page_break()
    sh = doc.add_paragraph()
    sr = sh.add_run("Sources")
    sr.bold, sr.font.size, sr.font.color.rgb = True, Pt(16), SLATE
    for name, url in SOURCES_SHORT:
        doc.add_paragraph(f"{name}: {url}", style="List Bullet")
    doc.add_paragraph(
        "Sentence-by-sentence source mapping is in FACT_CHECK.md; full notes in SOURCES.md. "
        "All images were drawn from scratch for this entry (no stock art, logos or AI images). "
        "Font: Readex Pro, SIL Open Font License.")
    path = ENTRY / "Redbelly_Thread.docx"
    doc.save(path)
    print("wrote", path.relative_to(ROOT))
    return path


def build_pdf(docx_path):
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        print("LibreOffice not found; skipped PDF")
        return
    subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(ENTRY),
                    str(docx_path)], check=True, capture_output=True)
    print("wrote entry/Redbelly_Thread.pdf")


if __name__ == "__main__":
    build_markdown()
    build_pdf(build_docx())
