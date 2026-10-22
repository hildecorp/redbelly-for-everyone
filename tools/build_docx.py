"""Builds entry/Redbelly_Thread.docx in the DAO brand, with the brand fonts embedded.

Be Vietnam Pro (regular and bold) and JetBrains Mono are embedded as obfuscated
TrueType fonts (ECMA-376 Part 1, 17.8.1), so the document looks the same on a
computer that does not have them installed.
"""

import io
import re
from datetime import datetime
import sys
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from build import CHECKLIST, N, N_IMG, fmt_count  # noqa: E402
from content import (  # noqa: E402
    FACTS, LAST_VERIFIED, LAST_VERIFIED_LONG, LEFT_OUT, POSTS, SOURCES, SUBTITLE, TASK, TITLE,
    post_sources,
)

SANS, MONO = "Be Vietnam Pro", "JetBrains Mono"
INK, INK2, MUTED = "0F1A22", "394853", "596975"
ACCENT, BRAND, OK = "B3272A", "EF5350", "15803D"
STEEL, LINE, LINE2, WELL = "1D3140", "D3DBE1", "DAE1E7", "F4F6F8"
TEXT_W = Cm(17)


# ---------------------------------------------------------------- xml helpers

def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def shade(cell_or_par, fill):
    target = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else cell_or_par._p.get_or_add_pPr()
    target.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def cell_borders(cell, color=LINE, sz=6, sides=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    b = _el("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            b.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": sz, "w:space": 0, "w:color": color}))
        else:
            b.append(_el(f"w:{side}", **{"w:val": "nil"}))
    tcPr.append(b)


def cell_margins(cell, top=200, left=240, bottom=200, right=240):
    tcPr = cell._tc.get_or_add_tcPr()
    m = _el("w:tcMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        m.append(_el(f"w:{side}", **{"w:w": v, "w:type": "dxa"}))
    tcPr.append(m)


def no_split(row):
    row._tr.get_or_add_trPr().append(_el("w:cantSplit"))


def par_border(par, side="left", color=BRAND, sz=18, space=8):
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = _el("w:pBdr")
        pPr.append(bdr)
    bdr.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": sz, "w:space": space, "w:color": color}))


def table_width(table, width=TEXT_W):
    tblPr = table._tbl.tblPr
    tblPr.append(_el("w:tblW", **{"w:w": int(width.twips), "w:type": "dxa"}))
    tblPr.append(_el("w:tblLayout", **{"w:type": "fixed"}))


def col_widths(table, widths):
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = w


def run(par, text, size=10.5, color=INK, bold=False, mono=False, caps=False, spacing=None, italic=False):
    r = par.add_run(text)
    f = r.font
    f.name = MONO if mono else SANS
    r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), MONO if mono else SANS)
    f.size = Pt(size)
    f.color.rgb = RGBColor.from_string(color)
    f.bold = bold
    f.italic = italic
    f.all_caps = caps
    if spacing is not None:
        r._r.get_or_add_rPr().append(_el("w:spacing", **{"w:val": spacing}))
    return r


def para(container, before=0, after=4, line=1.35, keep_next=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    pf.keep_with_next = keep_next
    return p


def field(par, code, size=7, color=MUTED):
    r = run(par, "", size=size, color=color, mono=True)
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    instr = _el("w:instrText", **{"xml:space": "preserve"})
    instr.text = f" {code} "
    r2 = run(par, "", size=size, color=color, mono=True)
    r2._r.append(instr)
    r3 = run(par, "", size=size, color=color, mono=True)
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    run(par, "1", size=size, color=color, mono=True)
    r5 = run(par, "", size=size, color=color, mono=True)
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def kicker(doc, text):
    p = para(doc, before=0, after=2, keep_next=True)
    run(p, text, size=7.5, color=ACCENT, bold=True, caps=True, spacing=20)


def heading(doc, text, intro=None):
    p = para(doc, after=3, keep_next=True)
    run(p, text, size=17, bold=True)
    if intro:
        p = para(doc, after=10, keep_next=True)
        run(p, intro, size=10.5, color=INK2)


def one_cell(doc, fill=None, border=LINE, margins=(200, 260, 200, 260)):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width(t)
    c = t.cell(0, 0)
    c.width = TEXT_W
    if fill:
        shade(c, fill)
    if border:
        cell_borders(c, border)
    else:
        cell_borders(c, sides=())
    cell_margins(c, *margins)
    no_split(t.rows[0])
    return t, c


# ---------------------------------------------------------------- document

def build_document():
    doc = Document()
    cp = doc.core_properties
    cp.title, cp.subject, cp.author, cp.last_modified_by, cp.comments, cp.keywords = TITLE, SUBTITLE, "", "", "", ""
    cp.created = cp.modified = datetime.strptime(LAST_VERIFIED, "%d %b %Y")

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
    sec.footer_distance = Cm(0.9)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = SANS, Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    rf = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), SANS)

    # Footer: title, task, date, page x of y.
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    run(fp, f"{TITLE}   {TASK}   Last verified {LAST_VERIFIED}", size=7, color=MUTED, mono=True, caps=True)
    run(fp, "\tPage ", size=7, color=MUTED, mono=True, caps=True)
    field(fp, "PAGE")
    run(fp, " of ", size=7, color=MUTED, mono=True, caps=True)
    field(fp, "NUMPAGES")

    # Cover band.
    _, c = one_cell(doc, fill=STEEL, border=None, margins=(420, 440, 420, 440))
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(14)
    p.add_run().add_picture(str(ROOT / "assets" / "brand" / "dao-logo-on-dark.png"), height=Cm(1.6))
    p = para(c, after=4)
    run(p, f"Redbelly DAO Task Board, {TASK}", size=7.5, color="8FA1AD", bold=True, caps=True, spacing=20)
    p = para(c, after=6, line=1.05)
    run(p, TITLE, size=28, color="F1F5F8", bold=True)
    p = para(c, after=10)
    run(p, SUBTITLE, size=12, color="C2CED6")
    p = para(c, after=0)
    run(p, f"Last verified {LAST_VERIFIED}     {N} posts     {N_IMG} images     ", size=7.5, color="C2CED6", mono=True, caps=True)
    run(p, "● ", size=7.5, color="86EFAC", mono=True)
    run(p, "All within 280 characters", size=7.5, color="C2CED6", mono=True, caps=True)

    p = para(doc, before=14, after=14, line=1.45)
    shade(p, WELL)
    par_border(p, "left", BRAND, 18, 8)
    run(p, "How to post it. ", bold=True)
    run(p, f"Post the {N} posts in order as one thread on X. Attach each image to the post it sits under, "
           "and paste its alt text into the image description. Do not post before judging ends: the DAO "
           "posts the winning entry.", color=INK2)

    kicker(doc, "At a glance")
    heading(doc, f"The thread in {N} posts")
    t = doc.add_table(rows=0, cols=4)
    table_width(t)
    gw = (Cm(1.4), Cm(10.6), Cm(2.4), Cm(2.6))
    for i, post in enumerate(POSTS, 1):
        row = t.add_row()
        for c in row.cells:
            cell_borders(c, LINE2, 4, ("bottom",))
            cell_margins(c, 70, 60, 70, 60)
        run(row.cells[0].paragraphs[0], f"{i:02d}", size=9, color=ACCENT, bold=True, mono=True)
        run(row.cells[1].paragraphs[0], post["title"], size=9.5)
        if post.get("image"):
            run(row.cells[2].paragraphs[0], f"Image {post['image'][:2]}", size=8, color=MUTED, mono=True)
        rp = row.cells[3].paragraphs[0]
        rp.alignment = 2
        run(rp, fmt_count(post), size=8, color=MUTED, mono=True)
    col_widths(t, gw)
    p = para(doc, before=12, after=0)
    run(p, "In this document: ", size=10, bold=True)
    run(p, "the thread, the brief checklist, the sources (deliverable 2) and a sentence-by-sentence fact check.", size=10, color=INK2)
    p.add_run().add_break(WD_BREAK.PAGE)

    kicker(doc, "The entry")
    heading(doc, "The thread", "Written for a friend, a parent or a business owner. Source IDs under each "
                               "post point to the source list.")

    inner_w = TEXT_W - Cm(0.92)
    for i, post in enumerate(POSTS, 1):
        _, c = one_cell(doc, fill="FFFFFF", border=LINE)
        p = c.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(inner_w, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        par_border(p, "bottom", LINE2, 6, 4)
        run(p, f"{i:02d} / {N}   ", size=9, color=ACCENT, bold=True, mono=True)
        run(p, post["title"], size=11, bold=True)
        run(p, f"\t{fmt_count(post)}", size=8, color=MUTED, mono=True)
        for chunk in post["text"].split("\n\n"):
            p = para(c, before=4, after=2, line=1.4)
            run(p, chunk, size=11)
        if post.get("image"):
            p = para(c, before=6, after=4)
            p.add_run().add_picture(str(ROOT / "entry" / "images" / post["image"]), width=inner_w)
            p = para(c, after=4)
            run(p, "Alt text  ", size=7.5, color=INK2, bold=True, caps=True, spacing=16)
            run(p, post["alt"], size=8.5, color=MUTED)
        p = para(c, before=2, after=0)
        run(p, "Sources: " + ", ".join(post_sources(post)), size=7.5, color=MUTED, mono=True)
        para(doc, after=4)

    # Checklist.
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    kicker(doc, "For the reviewer")
    heading(doc, "Brief checklist", "Each failure condition in the brief, and how this entry avoids it.")
    t = doc.add_table(rows=1, cols=3)
    table_width(t)
    widths = (Cm(4.2), Cm(2.2), Cm(10.6))
    col_widths(t, widths)
    for j, h in enumerate(("Rule", "Status", "How")):
        c = t.rows[0].cells[j]
        c.width = widths[j]
        shade(c, WELL)
        cell_borders(c, LINE, 6, ("bottom",))
        cell_margins(c, 100, 120, 100, 120)
        run(c.paragraphs[0], h, size=7, color=MUTED, bold=True, caps=True, spacing=20)
    for rule, how, verify in CHECKLIST:
        row = t.add_row()
        no_split(row)
        for j in range(3):
            row.cells[j].width = widths[j]
            cell_borders(row.cells[j], LINE2, 4, ("bottom",))
            cell_margins(row.cells[j], 110, 120, 110, 120)
        run(row.cells[0].paragraphs[0], rule, size=9.5, bold=True)
        sp = row.cells[1].paragraphs[0]
        run(sp, "● ", size=8, color=OK, mono=True)
        run(sp, "PASS", size=8, color=OK, bold=True, mono=True)
        run(row.cells[2].paragraphs[0], how, size=9.5, color=INK2)
        run(para(row.cells[2], before=2, after=0), verify, size=8.5, color=MUTED)
    col_widths(t, widths)

    # Sources.
    p = para(doc, before=18, after=0)
    kicker(doc, "Deliverable 2")
    heading(doc, "Sources", f"Every fact in the thread comes from one of these public sources, checked on {LAST_VERIFIED_LONG}.")
    t = doc.add_table(rows=0, cols=2)
    table_width(t)
    for s in SOURCES:
        row = t.add_row()
        no_split(row)
        a, b = row.cells
        a.width, b.width = Cm(1.4), Cm(15.6)
        for c in (a, b):
            cell_borders(c, LINE2, 4, ("top",))
            cell_margins(c, 110, 60, 110, 60)
        run(a.paragraphs[0], s["id"], size=9, color=ACCENT, bold=True, mono=True)
        p = b.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        run(p, s["name"], size=10, bold=True)
        run(para(b, after=1), s["publisher"], size=8.5, color=MUTED)
        run(para(b, after=2), s["url"], size=7.5, color=ACCENT, mono=True)
        run(para(b, after=0), s["used_for"], size=9, color=INK2)
    col_widths(t, (Cm(1.4), Cm(15.6)))

    # Fact check.
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    kicker(doc, "Verification")
    heading(doc, "Fact check", "Each claim the thread makes, with the exact source wording behind it. "
                               "\"[...]\" marks a cut. S6 page numbers are the report's printed pages.")
    for f in FACTS:
        used = ", ".join(str(i) for i, p_ in enumerate(POSTS, 1) if f["id"] in p_["facts"])
        p = para(doc, before=8, after=3, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
        par_border(p, "top", LINE2, 6, 6)
        run(p, f"{f['id']}  ", size=9.5, color=ACCENT, bold=True, mono=True)
        run(p, f["label"], size=10.5, bold=True)
        run(p, f"\tPost {used}", size=7.5, color=MUTED, mono=True)
        run(para(doc, after=4, keep_next=True), f["claim"], size=10)
        for sid, q in f["evidence"]:
            p = para(doc, after=3, line=1.35)
            p.paragraph_format.left_indent = Cm(0.3)
            shade(p, WELL)
            par_border(p, "left", "C3CDD4", 12, 6)
            run(p, f"{sid}  ", size=8, color=ACCENT, bold=True, mono=True)
            run(p, q, size=9, color=INK2)
        if f.get("note"):
            run(para(doc, before=1, after=2), f["note"], size=8.5, color=MUTED, italic=True)
    p = para(doc, before=12)
    run(p, "Left out on purpose: ", size=9.5, bold=True)
    run(p, " ".join(LEFT_OUT), size=9.5, color=INK2)
    return doc


# ---------------------------------------------------------------- font embedding

def _ttf_bytes():
    """Static TrueType copies of the web fonts, named the way the document refers to them."""
    from fontTools.ttLib import TTFont
    from fontTools.varLib.instancer import instantiateVariableFont
    fonts = ROOT / "assets" / "fonts"

    def rename(font, family, style):
        name = font["name"]
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, family), (2, style), (4, f"{family} {style}"),
                         (6, f"{family.replace(' ', '')}-{style}")):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        font["OS/2"].fsType = 0
        out = io.BytesIO()
        font.flavor = None
        font.save(out)
        return out.getvalue()

    reg = TTFont(str(fonts / "BeVietnamPro-400.woff2"))
    bold = TTFont(str(fonts / "BeVietnamPro-700.woff2"))
    mono = instantiateVariableFont(TTFont(str(fonts / "JetBrainsMono-var.woff2")), {"wght": 500})
    mono_bold = instantiateVariableFont(TTFont(str(fonts / "JetBrainsMono-var.woff2")), {"wght": 700})
    return {
        (SANS, "Regular"): rename(reg, SANS, "Regular"),
        (SANS, "Bold"): rename(bold, SANS, "Bold"),
        (MONO, "Regular"): rename(mono, MONO, "Regular"),
        (MONO, "Bold"): rename(mono_bold, MONO, "Bold"),
    }


def _obfuscate(data, guid):
    key = bytes.fromhex(guid.replace("-", ""))[::-1]
    head = bytes(b ^ key[i % 16] for i, b in enumerate(data[:32]))
    return head + data[32:]


def embed_fonts(docx_bytes):
    fonts = _ttf_bytes()
    zin = zipfile.ZipFile(io.BytesIO(docx_bytes))
    files = {n: zin.read(n) for n in zin.namelist()}

    rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    entries = {}
    for k, ((family, style), data) in enumerate(fonts.items(), 1):
        guid = str(uuid.uuid4()).upper()
        files[f"word/fonts/font{k}.odttf"] = _obfuscate(data, guid)
        rels.append(f'<Relationship Id="rIdFont{k}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/font{k}.odttf"/>')
        tag = "w:embedRegular" if style == "Regular" else "w:embedBold"
        entries.setdefault(family, []).append(f'<{tag} r:id="rIdFont{k}" w:fontKey="{{{guid}}}"/>')
    rels.append("</Relationships>")
    files["word/_rels/fontTable.xml.rels"] = "".join(rels).encode()

    ft = files["word/fontTable.xml"].decode("utf8")
    if 'xmlns:r="' not in ft.split(">", 2)[1]:
        ft = re.sub(r"<w:fonts ", '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', ft, count=1)
    for family, embeds in entries.items():
        block = (f'<w:font w:name="{family}"><w:charset w:val="00"/><w:family w:val="{"modern" if family == MONO else "swiss"}"/>'
                 f'<w:pitch w:val="{"fixed" if family == MONO else "variable"}"/>{"".join(embeds)}</w:font>')
        ft = ft.replace("</w:fonts>", block + "</w:fonts>")
    files["word/fontTable.xml"] = ft.encode("utf8")

    st = files["word/settings.xml"].decode("utf8")
    st = re.sub(r"(<w:settings[^>]*>)", r"\1<w:embedTrueTypeFonts/>", st, count=1)
    files["word/settings.xml"] = st.encode("utf8")

    ct = files["[Content_Types].xml"].decode("utf8")
    if 'Extension="odttf"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ', 1)
    files["[Content_Types].xml"] = ct.encode("utf8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", files.pop("[Content_Types].xml"))
        for n, d in files.items():
            z.writestr(n, d)
    return out.getvalue()


def build():
    buf = io.BytesIO()
    build_document().save(buf)
    path = ROOT / "entry" / "Redbelly_Thread.docx"
    path.write_bytes(embed_fonts(buf.getvalue()))
    print("wrote", path.relative_to(ROOT))


if __name__ == "__main__":
    build()
